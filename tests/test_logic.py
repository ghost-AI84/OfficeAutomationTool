# -*- coding: utf-8 -*-
"""合并引擎测试：顺序、表头对齐、类型、取消、原子写入。

不依赖界面，直接 `python tests/test_logic.py` 运行。
"""
import sys
import tempfile
import threading
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import pandas as pd

from excel_merge.logic import MergeCancelled, MergeOptions, merge_files

WORK = Path(tempfile.mkdtemp(prefix="merge-test-"))


def make(name, headers, rows):
    path = WORK / name
    pd.DataFrame(rows, columns=headers).to_excel(path, index=False)
    return str(path)


def merge_groups(files, **kw):
    """返回完整结果：工作表名 -> DataFrame。"""
    opts = MergeOptions(output_path=str(WORK / kw.pop("out", "out.xlsx")), **kw)
    logs = []
    result = merge_files(files, opts, log_callback=logs.append)
    return result, logs


def run(files, **kw):
    """只取第一张（单表模式下就是唯一一张）用于断言。"""
    result, logs = merge_groups(files, **kw)
    return next(iter(result.values())), logs


def test_order_defines_master_header():
    """用户列表顺序决定主表头，即使文件名排序相反。"""
    f1 = make("z_first.xlsx", ["姓名", "身份证号码", "金额(元)"],
              [["张三", "110101199001011234", "1,200.50"]])
    f2 = make("a_second.xlsx", [" 姓名 ", "证件号码", "金额"],
              [["李四", "310101199002022345", "800"]])
    f3 = make("m_third.xlsx", ["姓名", "备注"], [["王五", "老客户"]])

    df, _ = run([f1, f2, f3], out="ordered.xlsx")
    assert list(df.columns)[:3] == ["姓名", "身份证号码", "金额(元)"], list(df.columns)
    assert "备注" in df.columns
    assert df["身份证号码"].notna().sum() == 2
    assert list(df["数据来源"]) == ["z_first", "a_second", "m_third"]
    print("  order + synonym + extra column   OK")


def test_sort_by_name():
    f1 = make("z_first2.xlsx", ["姓名", "身份证号码"], [["张三", "1"]])
    f2 = make("a_second2.xlsx", ["姓名", "证件号码"], [["李四", "2"]])
    df, _ = run([f1, f2], out="sorted.xlsx", sort_by_name=True)
    assert "证件号码" in df.columns and "身份证号码" not in df.columns, list(df.columns)
    print("  sort_by_name flips master header  OK")


def test_cancel_leaves_no_output():
    f = make("cancel.xlsx", ["姓名"], [["张三"]])
    ev = threading.Event()
    ev.set()
    out = WORK / "cancelled.xlsx"
    try:
        merge_files([f], MergeOptions(output_path=str(out)), cancel_event=ev)
        raise AssertionError("cancel was not honoured")
    except MergeCancelled:
        pass
    assert not out.exists()
    assert not list(WORK.glob("cancelled*.part"))
    assert not list(WORK.glob("cancelled*tmp*"))
    print("  cancel + no leftovers             OK")


def test_csv_and_source_first():
    f = make("csv.xlsx", ["姓名", "金额"], [["张三", "100"]])
    df, _ = run([f], out="out.csv", source_position="first")
    assert list(df.columns)[0] == "数据来源", list(df.columns)
    assert (WORK / "out.csv").exists()
    assert not list(WORK.glob("out.csv.*"))
    print("  csv output + source first         OK")


def test_force_types():
    f = make("types.xlsx", ["金额(元)", "身份证号码"],
             [["1,200.50", "110101199001011234"]])
    df, _ = run([f], out="types_out.xlsx",
                force_types={"金额(元)": "number", "身份证号码": "text"})
    assert str(df["金额(元)"].dtype) in ("Float64", "float64"), df["金额(元)"].dtype
    assert df["金额(元)"].iloc[0] == 1200.5, df["金额(元)"].iloc[0]
    print("  force column types                OK")


def test_auto_header_detection():
    path = WORK / "with_title.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame([["2026 年销售报表", None, None],
                      [None, None, None],
                      ["姓名", "手机号", "金额"],
                      ["赵六", "13800000000", "300"]]).to_excel(w, index=False, header=False)
    df, logs = run([str(path)], out="detected.xlsx", header_row=None)
    assert list(df.columns)[:3] == ["姓名", "手机号", "金额"], list(df.columns)
    assert any("自动检测表头" in m for m in logs)
    print("  auto header detection             OK")


def test_drop_empty_rows_and_cols():
    f = make("blanks.xlsx", ["姓名", "空列", "金额"],
             [["钱七", None, "100"], [None, None, None], ["孙八", None, "200"]])
    df, _ = run([f], out="blanks_out.xlsx")
    assert len(df) == 2 and "空列" in df.columns
    df2, _ = run([f], out="blanks_out2.xlsx", drop_empty_cols=True)
    assert "空列" not in df2.columns, list(df2.columns)
    print("  drop empty rows / columns         OK")


def test_source_column_name_clash():
    f = make("clash.xlsx", ["姓名", "数据来源"], [["周九", "原值"]])
    df, _ = run([f], out="clash_out.xlsx")
    assert list(df["数据来源"]) == ["clash"], df["数据来源"].tolist()
    print("  source column name clash          OK")


def _write_grid(name, rows):
    """按原样写二维表（不设表头），返回路径。"""
    path = WORK / name
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame(rows).to_excel(w, index=False, header=False)
    return str(path)


def test_headers_at_different_rows():
    """每个文件表头行不同：一个有标题有空格，一个直接是表头。"""
    a = make("hdr_a.xlsx", ["姓名", "金额"], [["张三", "100"]])
    b = _write_grid("hdr_b.xlsx", [
        ["销售明细", None],
        [None, None],
        ["姓名", "金额"],
        ["李四", "200"],
    ])
    df, logs = run([a, b], out="hdr_diff.xlsx", header_row=None)
    assert list(df.columns)[:2] == ["姓名", "金额"], list(df.columns)
    assert len(df) == 2, len(df)
    assert any("自动检测表头" in m for m in logs)
    print("  headers at different rows         OK")


def test_per_file_header_override():
    """逐文件指定表头行，优先于自动识别；未指定的文件退回自动检测。"""
    b = _write_grid("hdr_override.xlsx", [
        ["销售明细", None],
        [None, None],
        ["姓名", "金额"],
        ["李四", "200"],
    ])
    df, _ = run([b], out="hdr_override_out.xlsx", header_row={b: 1})
    assert str(df.columns[0]) == "销售明细", list(df.columns)

    # 字典里没有这个文件 -> 仍然自动检测到第 3 行
    c = _write_grid("hdr_unlisted.xlsx", [
        ["说明文字", None],
        [None, None],
        ["姓名", "金额"],
        ["王五", "300"],
    ])
    df2, logs = run([c], out="hdr_unlisted_out.xlsx", header_row={b: 1})
    assert list(df2.columns)[:2] == ["姓名", "金额"], list(df2.columns)
    assert any("未单独指定" in m for m in logs)
    print("  per-file header override          OK")


def test_two_row_header_with_merged_cells():
    """两级表头：分组名是合并单元格，下一行才是子表头。"""
    from openpyxl import Workbook
    path = WORK / "two_row.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "姓名"
    ws["B1"] = "上半年"
    ws.merge_cells("B1:C1")
    ws["D1"] = "下半年"
    ws.merge_cells("D1:E1")
    for col, text in enumerate(["", "1月", "2月", "3月", "4月"], start=1):
        if text:
            ws.cell(row=2, column=col, value=text)
    ws.append(["张三", "10", "20", "30", "40"])
    ws.append(["李四", "11", "21", "31", "41"])
    wb.save(path)

    df, logs = run([str(path)], out="two_row_out.xlsx", header_row=None)
    assert list(df.columns)[:5] == ["姓名", "上半年_1月", "上半年_2月",
                                    "下半年_3月", "下半年_4月"], list(df.columns)
    assert len(df) == 2, len(df)
    assert "2-2" not in str(logs)
    print("  two-row header + merged cells     OK")


def test_drop_repeated_header_row():
    """分页导出：数据中间又出现一次表头。"""
    path = _write_grid("repeat_hdr.xlsx", [
        ["姓名", "金额"],
        ["张三", "100"],
        ["姓名", "金额"],
        ["李四", "200"],
    ])
    df, logs = run([str(path)], out="repeat_hdr_out.xlsx", header_row=None)
    assert len(df) == 2, len(df)
    assert any("重复表头" in m for m in logs)
    print("  repeated header row dropped       OK")


def test_drop_summary_row():
    """末尾合计行不该进入结果。"""
    path = _write_grid("summary.xlsx", [
        ["姓名", "金额"],
        ["张三", "100"],
        ["李四", "200"],
        ["合计", "300"],
    ])
    df, logs = run([str(path)], out="summary_out.xlsx", header_row=None)
    assert len(df) == 2, len(df)
    assert any("合计" in m for m in logs)
    print("  summary row dropped               OK")


def test_sheet_auto():
    from openpyxl import Workbook
    path = WORK / "sheets.xlsx"
    wb = Workbook()
    note = wb.active
    note.title = "说明"
    note["A1"] = "本表仅供内部使用"
    data = wb.create_sheet("数据")
    data.append(["姓名", "金额"])
    data.append(["张三", "100"])
    wb.save(path)

    df, logs = run([str(path)], out="sheets_out.xlsx", header_row=None, sheet_name="auto")
    assert list(df.columns)[:2] == ["姓名", "金额"], list(df.columns)
    assert any("自动选择工作表" in m for m in logs)
    print("  sheet auto-pick                   OK")


def test_autofit_column_widths():
    """列宽自适应必须真的写进文件（曾经因为临时文件扩展名被静默跳过）。"""
    from openpyxl import load_workbook
    f = make("width.xlsx", ["姓名", "身份证号码", "备注"],
             [["张三", "110101199001011234", "一段比较长的备注文字内容"]])
    out = WORK / "width_out.xlsx"
    run([f], out="width_out.xlsx")
    ws = load_workbook(out).active
    widths = [ws.column_dimensions[c].width for c in ("A", "B", "C")]
    assert widths[1] > 13, widths      # 身份证号列应明显宽于默认 13
    assert widths[2] > 13, widths
    assert widths[0] >= 8, widths      # 最小宽度
    print("  autofit writes column widths      OK  %s" % [round(w, 1) for w in widths])


def test_all_sheets_in_one_workbook():
    """多个子表按名字分组：附件1 与附件1 合并，附件2 与附件2 合并，各成一个 sheet。"""
    from openpyxl import Workbook, load_workbook

    def build(filename, rows1, rows2):
        p = WORK / filename
        wb = Workbook()
        wb.remove(wb.active)
        ws1 = wb.create_sheet("附件1")
        ws1.append(["姓名", "金额"])
        for r in rows1:
            ws1.append(r)
        ws2 = wb.create_sheet("附件2")
        ws2.append(["单位", "数量"])
        for r in rows2:
            ws2.append(r)
        note = wb.create_sheet("说明")
        note["A1"] = "本表仅为说明"
        wb.save(p)
        return str(p)

    a = build("multi_a.xlsx", [["甲", "100"]], [["A局", "5"]])
    b = build("multi_b.xlsx", [["乙", "200"]], [["B局", "7"]])

    out = WORK / "multi_grouped.xlsx"
    groups, logs = merge_groups([a, b], out="multi_grouped.xlsx",
                                sheet_name="all", header_row=None)

    assert list(groups.keys()) == ["附件1", "附件2"], list(groups.keys())
    assert len(groups["附件1"]) == 2, len(groups["附件1"])
    assert len(groups["附件2"]) == 2, len(groups["附件2"])
    assert list(groups["附件1"].columns)[:2] == ["姓名", "金额"]
    assert list(groups["附件2"].columns)[:2] == ["单位", "数量"]
    assert any("跳过空工作表" in m for m in logs)

    wb = load_workbook(out)
    assert wb.sheetnames == ["附件1", "附件2"], wb.sheetnames
    print("  sheets grouped by name            OK  %s" % wb.sheetnames)

    # 默认（只读第一张）仍然可用，不破坏原行为
    df2, _ = run([a], out="multi_single.xlsx", header_row=None)
    assert len(df2) == 1, len(df2)
    print("  default single-sheet unchanged    OK")


def test_sheets_matched_by_similar_name():
    """子表名不完全一致但相似：附件1-台账 / 附件1：台账 要配成一组。"""
    from openpyxl import Workbook

    def build(filename, names, rows):
        p = WORK / filename
        wb = Workbook()
        wb.remove(wb.active)
        for name in names:
            ws = wb.create_sheet(name)
            ws.append(["姓名", "金额"])
            for r in rows:
                ws.append(r)
        wb.save(p)
        return str(p)

    a = build("sim_a.xlsx", ["附件1-台账", "附件2-发票"], [["甲", "100"]])
    b = build("sim_b.xlsx", ["附件1：台账", "附件2（发票）"], [["乙", "200"]])

    groups, logs = merge_groups([a, b], out="sim_out.xlsx",
                                sheet_name="all", header_row=None)
    assert list(groups.keys()) == ["附件1-台账", "附件2-发票"], list(groups.keys())
    assert len(groups["附件1-台账"]) == 2, len(groups["附件1-台账"])
    assert len(groups["附件2-发票"]) == 2, len(groups["附件2-发票"])
    assert any("对应到" in m for m in logs)
    print("  similar sheet names matched       OK")


def test_sheets_matched_by_position():
    """名字差异很大（1.1 / 2.1）时按顺序配对。"""
    from openpyxl import Workbook

    def build(filename, names):
        p = WORK / filename
        wb = Workbook()
        wb.remove(wb.active)
        for i, name in enumerate(names):
            ws = wb.create_sheet(name)
            ws.append(["姓名", "金额"])
            ws.append([filename + str(i), str(i)])
        wb.save(p)
        return str(p)

    a = build("pos_a.xlsx", ["1.1", "1.2"])
    b = build("pos_b.xlsx", ["2.1", "2.2"])

    groups, logs = merge_groups([a, b], out="pos_out.xlsx",
                                sheet_name="all", header_row=None)
    assert list(groups.keys()) == ["1.1", "1.2"], list(groups.keys())
    assert len(groups["1.1"]) == 2 and len(groups["1.2"]) == 2
    print("  sheets matched by position        OK")


def test_extra_sheet_not_force_matched():
    """子表数量对不上时不能硬配：多出来的「汇总」表要跳过，而不是塞进第 1 组。"""
    from openpyxl import Workbook

    def build(filename, sheets):
        p = WORK / filename
        wb = Workbook()
        wb.remove(wb.active)
        for name in sheets:
            ws = wb.create_sheet(name)
            ws.append(["姓名", "金额"])
            ws.append([filename + name, "1"])
        wb.save(p)
        return str(p)

    a = build("cnt_a.xlsx", ["附件1", "附件2"])
    b = build("cnt_old_合并.xlsx", ["汇总"])

    groups, logs = merge_groups([a, b], out="cnt_out.xlsx",
                                sheet_name="all", header_row=None)
    assert list(groups.keys()) == ["附件1", "附件2"], list(groups.keys())
    assert len(groups["附件1"]) == 1, len(groups["附件1"])   # 只有 a 的行
    assert any("没有对得上的组" in m for m in logs)
    print("  extra sheet not force-matched    OK")


def _inject_text_autofilter(path):
    """给工作表塞一个"文字型自定义筛选"，复现 openpyxl 3.1.2 的解析崩溃。"""
    import shutil
    import zipfile

    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                if "<autoFilter" not in text:
                    # 必须包在 filterColumn 里，否则 openpyxl 不会解析这段
                    flt = ('<autoFilter ref="A1:B3"><filterColumn colId="0">'
                           '<customFilters><customFilter operator="equal" val="张三"/>'
                           '</customFilters></filterColumn></autoFilter>')
                    text = text.replace("</sheetData>", "</sheetData>" + flt)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


def test_text_autofilter_workbook_is_readable():
    """带文字型自动筛选的表要能读——openpyxl 3.1.2 会直接 ValueError。"""
    path = WORK / "autofilter.xlsx"
    pd.DataFrame([["姓名", "金额"], ["张三", "100"], ["李四", "200"]]).to_excel(
        path, index=False, header=False)
    _inject_text_autofilter(path)

    df, _ = run([str(path)], out="autofilter_out.xlsx", header_row=None)
    assert list(df.columns)[:2] == ["姓名", "金额"], list(df.columns)
    assert len(df) == 2, len(df)
    print("  text autofilter workbook          OK")


TESTS = [
    test_order_defines_master_header,
    test_sort_by_name,
    test_cancel_leaves_no_output,
    test_csv_and_source_first,
    test_force_types,
    test_auto_header_detection,
    test_drop_empty_rows_and_cols,
    test_source_column_name_clash,
    test_headers_at_different_rows,
    test_per_file_header_override,
    test_two_row_header_with_merged_cells,
    test_drop_repeated_header_row,
    test_drop_summary_row,
    test_sheet_auto,
    test_autofit_column_widths,
    test_all_sheets_in_one_workbook,
    test_sheets_matched_by_similar_name,
    test_sheets_matched_by_position,
    test_extra_sheet_not_force_matched,
    test_text_autofilter_workbook_is_readable,
]


def main():
    print("logic tests (%d)" % len(TESTS))
    for fn in TESTS:
        fn()
    print("logic tests passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
