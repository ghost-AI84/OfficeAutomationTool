# -*- coding: utf-8 -*-
"""端到端测试：通过界面真正跑一次合并，验证线程、消息队列与输出文件。"""
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

os.environ["APPDATA"] = tempfile.mkdtemp(prefix="merge-appdata-e2e-")

import pandas as pd
import tkinter as tk
from tkinter import messagebox, ttk

from excel_merge.ui import MergeApp, enable_dpi_awareness, apply_default_fonts


def main():
    work = Path(tempfile.mkdtemp(prefix="merge-e2e-"))
    f1 = work / "b_file.xlsx"
    f2 = work / "a_file.xlsx"
    pd.DataFrame([["张三", "110101199001011234"]],
                 columns=["姓名", "身份证"]).to_excel(f1, index=False)
    pd.DataFrame([["李四", "310101199002022345"]],
                 columns=["姓名 ", "身份证号码"]).to_excel(f2, index=False)

    dialogs = []
    messagebox.askyesno = lambda *a, **k: (dialogs.append(("askyesno", a[0])), False)[1]
    messagebox.showwarning = lambda *a, **k: dialogs.append(("warn", a[1] if len(a) > 1 else a))
    messagebox.showerror = lambda *a, **k: dialogs.append(("error", a[1] if len(a) > 1 else a))
    messagebox.showinfo = lambda *a, **k: dialogs.append(("info", a[1] if len(a) > 1 else a))

    enable_dpi_awareness()
    try:
        root = tk.Tk()
    except tk.TclError as e:
        print("SKIP: no usable display (%s)" % e)
        return 0
    apply_default_fonts(root)
    try:
        ttk.Style().theme_use("clam")
    except tk.TclError:
        pass

    app = MergeApp(root)
    root.update_idletasks()

    # b_file 在列表首位，应决定主表头（列名是「身份证」而不是「身份证号码」）
    app._add_paths([str(f1), str(f2)])
    assert app.files == [str(f1.resolve()), str(f2.resolve())], app.files

    out = work / "merged.xlsx"
    app.v_output.set(str(out))
    app.v_header_mode.set("fixed")
    app.v_header_row.set("1")
    app._sync_header_entry()

    app.start_merge()

    deadline = time.time() + 60
    while time.time() < deadline:
        root.update()
        if app.worker is not None and not app.worker.is_alive():
            for _ in range(5):
                time.sleep(0.1)
                root.update()
            break

    assert out.exists(), "output file not produced"
    df = pd.read_excel(out)
    assert "身份证" in df.columns and "身份证号码" not in df.columns, list(df.columns)
    assert len(df) == 2, len(df)
    assert app.v_status.get() == "完成", app.v_status.get()
    assert str(app.run_btn.cget("state")) == "normal"
    assert str(app.cancel_btn.cget("state")) == "disabled"
    assert app.progress["value"] == app.progress["maximum"]
    assert not (work / "merged.xlsx.part").exists()
    print("  merge through UI                  OK  columns=%s" % list(df.columns))

    # 覆盖确认返回 False 时不应启动
    dialogs.clear()
    app.start_merge()
    root.update()
    assert any(d[0] == "askyesno" for d in dialogs), dialogs
    assert app.worker is None or not app.worker.is_alive()
    print("  overwrite prompt cancels safely   OK")

    root.destroy()
    run_grouped_case()
    print("e2e tests passed")
    return 0


def run_grouped_case():
    """多个子表按名字分组：附件1 归附件1，附件2 归附件2，输出两个 sheet。"""
    from openpyxl import Workbook, load_workbook

    work = Path(tempfile.mkdtemp(prefix="merge-grouped-"))

    def build(name, rows1, rows2):
        p = work / name
        wb = Workbook()
        wb.remove(wb.active)
        ws1 = wb.create_sheet("附件1")
        ws1.append(["姓名", "金额"])
        for r in rows1:
            ws1.append(r)
        ws2 = wb.create_sheet("附件2")
        ws2.append(["单位", "数量"])
        for r in rows2:
            ws2.append(r)
        note = wb.create_sheet("说明")
        note["A1"] = "说明页"
        wb.save(p)
        return str(p)

    a = build("甲所.xlsx", [["张三", "100"], ["李四", "200"]], [["A局", "5"]])
    b = build("乙所.xlsx", [["王五", "300"]], [["B局", "7"], ["C局", "9"]])

    try:
        root = tk.Tk()
    except tk.TclError as e:
        print("SKIP grouped case: no display (%s)" % e)
        return
    apply_default_fonts(root)
    try:
        ttk.Style().theme_use("clam")
    except tk.TclError:
        pass

    app = MergeApp(root)
    root.update_idletasks()
    app._add_paths([a, b])
    out = work / "结果.xlsx"
    app.v_output.set(str(out))
    app.v_sheet.set("所有子表（按名字分组）")
    app.start_merge()

    deadline = time.time() + 60
    while time.time() < deadline:
        root.update()
        if app.worker is not None and not app.worker.is_alive():
            for _ in range(6):
                time.sleep(0.1)
                root.update()
            break

    assert out.exists(), "grouped output not produced"
    wb = load_workbook(out)
    assert wb.sheetnames == ["附件1", "附件2"], wb.sheetnames
    assert wb["附件1"].max_row == 4, wb["附件1"].max_row    # 3 条数据 + 表头
    assert wb["附件2"].max_row == 4, wb["附件2"].max_row
    headers1 = [c.value for c in wb["附件1"][1]][:2]
    headers2 = [c.value for c in wb["附件2"][1]][:2]
    assert headers1 == ["姓名", "金额"], headers1
    assert headers2 == ["单位", "数量"], headers2
    print("  grouped sheets through UI         OK  %s" % wb.sheetnames)
    root.destroy()


if __name__ == "__main__":
    sys.exit(main())
